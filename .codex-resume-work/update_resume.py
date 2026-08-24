from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\company_agent\.codex-resume-work\source.docx")
OUTPUT = Path(r"C:\Users\lenovo\Desktop\董星晨\董星晨的个人简历_AI应用开发版.docx")

FONT_CN = "Microsoft YaHei"
FONT_EN = "Calibri"
BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(35, 35, 35)
GRAY = RGBColor(89, 89, 89)


def set_font(run, size, *, bold=False, color=DARK, east_asia=FONT_CN):
    run.font.name = FONT_EN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), east_asia)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def write_paragraph(
    paragraph,
    text,
    *,
    size=10.2,
    bold=False,
    color=DARK,
    before=0,
    after=2,
    line=1.05,
    left=0,
    first=0,
    keep_with_next=False,
):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.left_indent = Inches(left)
    paragraph.paragraph_format.first_line_indent = Inches(first)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text)
    set_font(run, size, bold=bold, color=color)
    return paragraph


def write_bullet(paragraph, text):
    write_paragraph(
        paragraph,
        text,
        size=9.8,
        after=1.5,
        line=1.03,
        left=0.18,
        first=-0.18,
    )
    paragraph.style = "List Bullet"


def write_project_title(paragraph, title):
    write_paragraph(
        paragraph,
        title,
        size=11.2,
        bold=True,
        after=2,
        line=1.0,
        keep_with_next=True,
    )


def write_cell(cell, text, *, size=9.5, bold=False, color=DARK):
    paragraph = cell.paragraphs[0]
    write_paragraph(paragraph, text, size=size, bold=bold, color=color, after=0, line=1.0)
    paragraph.paragraph_format.keep_together = True
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


doc = Document(SOURCE)

# Header: keep the existing photo in the right cell and replace only the text block.
header_cell = doc.tables[0].cell(0, 0)
while len(header_cell.paragraphs) < 3:
    header_cell.add_paragraph()

first = header_cell.paragraphs[0]
write_paragraph(first, "董星晨", size=22, bold=True, color=BLUE, after=4, line=1.0)

contact = header_cell.paragraphs[1]
write_paragraph(
    contact,
    "电话：18656018213  |  邮箱：d18656018213@outlook.com  |  所在地：合肥  |  到岗时间：三天内",
    size=9.4,
    color=GRAY,
    after=2,
    line=1.0,
)

target = header_cell.paragraphs[2]
write_paragraph(
    target,
    "求职方向：AI 应用开发实习生 / Python 后端实习生",
    size=9.6,
    bold=True,
    color=GRAY,
    after=0,
    line=1.0,
)

# Remove the empty/legacy header paragraphs left by the template.
for paragraph in list(header_cell.paragraphs)[3:]:
    paragraph._element.getparent().remove(paragraph._element)

paragraphs = doc.paragraphs

# Education.
write_paragraph(paragraphs[1], "安徽建筑大学  |  本科在读", size=10.4, bold=True, after=4)

# Primary project: actual implementation in D:\company_agent.
write_project_title(paragraphs[3], "企业制度知识库问答 Agent")
write_bullet(
    paragraphs[4],
    "技术栈：Python、FastAPI、LangChain Agent、DeepSeek、Ollama Embeddings、Chroma、Redis、HTML/CSS/JavaScript。",
)
write_bullet(
    paragraphs[5],
    "完成文档加载、递归切分、向量化与 Chroma 持久化检索，并将知识库检索封装为工具，由 Agent 按问题自主调用。",
)
write_bullet(
    paragraphs[6],
    "使用 FastAPI StreamingResponse 实现流式回答；通过 Cookie 自动传递会话标识，使用 Redis 保存最近 10 轮完整对话。",
)
write_bullet(
    paragraphs[7],
    "搭建响应式对话页面与新建会话功能；补充请求 ID、检索耗时、工具调用、历史读写及异常信息等日志链路。",
)

# Secondary project: retain only claims already present in the source resume.
write_project_title(paragraphs[8], "基于 CSV 岗位数据的 AI 求职问答")
write_bullet(
    paragraphs[9],
    "技术栈：Python、Pandas、LangChain、RAG、FastAPI、Streamlit。",
)
write_bullet(
    paragraphs[10],
    "面向岗位查询场景，对 CSV 岗位数据进行清洗与结构化处理，为条件筛选和问答检索提供统一数据来源。",
)
write_bullet(
    paragraphs[11],
    "结合 LangChain 与 RAG 完成岗位信息检索，使用户能够通过自然语言查询符合条件和预期的岗位。",
)
write_bullet(
    paragraphs[12],
    "使用 FastAPI 提供问答能力，并通过 Streamlit 搭建交互界面，完成从数据处理到结果展示的基础流程。",
)

# Skills table.
skill_rows = [
    (
        "Python / 后端",
        "掌握 Python 基础语法、面向对象、模块化与异常处理；能够使用 FastAPI、Pydantic 编写接口和流式响应。",
    ),
    (
        "大模型应用",
        "能够使用 LangChain 构建 Agent、工具调用与 RAG 流程，理解 Prompt、对话历史和检索结果的传递方式。",
    ),
    (
        "数据与存储",
        "使用 Pandas 处理 CSV 数据；使用 Chroma 完成向量存储与相似度检索，使用 Redis 管理会话历史。",
    ),
    (
        "前端与工程",
        "具备 HTML、CSS、JavaScript 基础，可完成响应式对话界面；能够配置依赖并使用日志定位请求链路问题。",
    ),
]

skills_table = doc.tables[1]
for row, (label, detail) in zip(skills_table.rows, skill_rows):
    write_cell(row.cells[0], label, size=9.3, bold=True, color=BLUE)
    write_cell(row.cells[1], detail, size=9.1)

# Compact cell padding and preserve the template's two-column visual structure.
for row in skills_table.rows:
    for cell in row.cells:
        tc_pr = cell._tc.get_or_add_tcPr()
        tc_mar = tc_pr.first_child_found_in("w:tcMar")
        if tc_mar is None:
            from docx.oxml import OxmlElement

            tc_mar = OxmlElement("w:tcMar")
            tc_pr.append(tc_mar)
        for edge, value in (("top", 55), ("bottom", 55), ("start", 90), ("end", 90)):
            node = tc_mar.find(qn(f"w:{edge}"))
            if node is None:
                from docx.oxml import OxmlElement

                node = OxmlElement(f"w:{edge}")
                tc_mar.append(node)
            node.set(qn("w:w"), str(value))
            node.set(qn("w:type"), "dxa")

# Self-summary: evidence-based and aligned with an entry-level AI application role.
write_paragraph(
    paragraphs[15],
    "能够独立完成一个小型 AI 应用从文档处理、向量检索、Agent 编排、接口开发到前端展示和日志排查的完整流程。重视问题定位与边界验证，计划继续学习 LangGraph，并在实际项目中提升测试、部署与工程化能力。",
    size=10.0,
    after=0,
    line=1.1,
)

# Footer role label.
for section in doc.sections:
    for paragraph in section.footer.paragraphs:
        if paragraph.text.strip():
            write_paragraph(
                paragraph,
                "董星晨 | AI 应用开发实习生",
                size=8.0,
                color=GRAY,
                after=0,
                line=1.0,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

# Keep A4 geometry and the compact margins from the original one-page template.
section = doc.sections[0]
section.top_margin = Inches(0.45)
section.bottom_margin = Inches(0.45)
section.left_margin = Inches(0.55)
section.right_margin = Inches(0.55)

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)

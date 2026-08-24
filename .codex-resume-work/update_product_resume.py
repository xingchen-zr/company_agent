from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path(r"D:\company_agent\.codex-resume-work\source.docx")
OUTPUT = Path(r"C:\Users\lenovo\Desktop\董星晨\董星晨的个人简历_产品方向版.docx")

FONT_CN = "Microsoft YaHei"
FONT_EN = "Calibri"
BLUE = RGBColor(31, 78, 121)
DARK = RGBColor(35, 35, 35)
GRAY = RGBColor(89, 89, 89)


def set_font(run, size, *, bold=False, color=DARK):
    run.font.name = FONT_EN
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), FONT_CN)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def clear_paragraph(paragraph):
    p = paragraph._element
    for child in list(p):
        if child.tag != qn("w:pPr"):
            p.remove(child)


def write_paragraph(
    paragraph,
    text,
    *,
    size=10.0,
    bold=False,
    color=DARK,
    before=0,
    after=2,
    line=1.04,
    left=0,
    first=0,
    keep_with_next=False,
):
    clear_paragraph(paragraph)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    paragraph.paragraph_format.line_spacing = line
    paragraph.paragraph_format.left_indent = Inches(left)
    paragraph.paragraph_format.first_line_indent = Inches(first)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    run = paragraph.add_run(text)
    set_font(run, size, bold=bold, color=color)
    return paragraph


def write_bullet(paragraph, text):
    write_paragraph(
        paragraph,
        text,
        size=9.7,
        after=1.5,
        line=1.03,
        left=0.18,
        first=-0.18,
    )
    paragraph.style = "List Bullet"


def write_project_title(paragraph, title):
    write_paragraph(
        paragraph,
        title,
        size=11.2,
        bold=True,
        after=2,
        line=1.0,
        keep_with_next=True,
    )


def write_cell(cell, text, *, size=9.2, bold=False, color=DARK):
    paragraph = cell.paragraphs[0]
    write_paragraph(paragraph, text, size=size, bold=bold, color=color, after=0, line=1.0)
    paragraph.paragraph_format.keep_together = True
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


doc = Document(SOURCE)

# Header: retain the original two-column table and portrait.
header_cell = doc.tables[0].cell(0, 0)
while len(header_cell.paragraphs) < 3:
    header_cell.add_paragraph()

write_paragraph(
    header_cell.paragraphs[0],
    "董星晨",
    size=22,
    bold=True,
    color=BLUE,
    after=4,
    line=1.0,
)
write_paragraph(
    header_cell.paragraphs[1],
    "电话：18656018213  |  邮箱：d18656018213@outlook.com  |  所在地：合肥  |  到岗时间：三天内",
    size=9.4,
    color=GRAY,
    after=2,
    line=1.0,
)
write_paragraph(
    header_cell.paragraphs[2],
    "求职方向：AI 产品实习生 / 产品助理 / 需求分析实习生",
    size=9.6,
    bold=True,
    color=GRAY,
    after=0,
    line=1.0,
)
for paragraph in list(header_cell.paragraphs)[3:]:
    paragraph._element.getparent().remove(paragraph._element)

paragraphs = doc.paragraphs
write_paragraph(paragraphs[1], "安徽建筑大学  |  本科在读", size=10.4, bold=True, after=4)

# Product-focused primary project.
write_project_title(paragraphs[3], "企业制度知识库问答 Agent")
write_bullet(
    paragraphs[4],
    "项目目标：面向企业员工查询制度、考核、福利和操作流程的场景，设计对话式知识查询工具，减少人工翻找文档的步骤。",
)
write_bullet(
    paragraphs[5],
    "功能设计：梳理文档入库、知识检索、AI 回答、历史对话和新建会话等核心流程，并明确无相关资料、会话中断等异常边界。",
)
write_bullet(
    paragraphs[6],
    "原型落地：结合 FastAPI、LangChain、Chroma 和 Redis 完成功能验证，使用 HTML/CSS/JavaScript 搭建响应式对话界面。",
)
write_bullet(
    paragraphs[7],
    "迭代排查：针对检索工具未触发、上下文传递、页面滚动和接口连接等问题逐层定位，并补充请求与异常日志方便复盘。",
)

# Product-focused secondary project.
write_project_title(paragraphs[8], "基于 CSV 岗位数据的 AI 求职问答")
write_bullet(
    paragraphs[9],
    "用户场景：针对岗位信息分散、人工筛选条件繁琐的问题，设计支持自然语言查询的岗位信息问答工具。",
)
write_bullet(
    paragraphs[10],
    "需求拆解：将用户目标拆分为岗位条件筛选、相关岗位检索和求职建议三个环节，明确 CSV 数据作为当前信息来源。",
)
write_bullet(
    paragraphs[11],
    "数据处理：使用 Pandas 完成岗位数据清洗与结构化，并结合 LangChain、RAG 实现自然语言查询与结果返回。",
)
write_bullet(
    paragraphs[12],
    "交互验证：通过 FastAPI 提供问答能力，使用 Streamlit 搭建演示界面，验证从用户提问到岗位结果展示的完整流程。",
)

# Skills aligned to entry-level product roles without claiming unsupported tools.
skill_rows = [
    (
        "需求与流程",
        "能够围绕用户场景拆解需求、梳理核心功能流程、识别异常边界，并通过功能清单和验证步骤推进实现。",
    ),
    (
        "AI 产品认知",
        "理解 RAG、Agent、Prompt、工具调用、向量检索和对话历史等基本机制，能够判断常见功能的技术实现边界。",
    ),
    (
        "数据与文档",
        "可使用 Pandas、Excel 进行基础数据整理与分析，使用 Word、PPT 完成需求说明、资料归纳和汇报材料。",
    ),
    (
        "原型与技术",
        "具备 Python、FastAPI、HTML、CSS、JavaScript 基础，可独立完成小型交互原型并结合日志定位问题。",
    ),
]

skills_table = doc.tables[1]
for row, (label, detail) in zip(skills_table.rows, skill_rows):
    write_cell(row.cells[0], label, size=9.3, bold=True, color=BLUE)
    write_cell(row.cells[1], detail, size=9.1)

write_paragraph(
    paragraphs[15],
    "具备将业务问题拆解为可实现功能的意识，能够独立完成小型 AI 产品从场景分析、流程设计到开发验证和问题复盘的完整过程。既关注用户操作体验，也能理解接口、数据与模型能力边界，希望从 AI 产品或产品助理岗位积累真实业务经验。",
    size=9.8,
    after=0,
    line=1.08,
)

for section in doc.sections:
    for paragraph in section.footer.paragraphs:
        if paragraph.text.strip():
            write_paragraph(
                paragraph,
                "董星晨 | AI 产品实习生",
                size=8.0,
                color=GRAY,
                after=0,
                line=1.0,
            )
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

OUTPUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUTPUT)
print(OUTPUT)
